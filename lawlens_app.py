import streamlit as st
import PyPDF2
import docx
import os
import requests
import json
from dotenv import load_dotenv
from googletrans import Translator

# Load environment variables
load_dotenv()

# Page configuration
st.set_page_config(
    page_title="LawLens - Your Legal Companion",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Enhanced Custom CSS with better styling
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');
    
    * {
        font-family: 'Inter', sans-serif;
    }
    
    .main-header {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        font-size: 4.2rem;
        font-weight: 800;
        text-align: center;
        margin-bottom: 0.5rem;
        padding: 1rem;
        text-shadow: 0 4px 8px rgba(0,0,0,0.1);
    }
    
    .tagline {
        text-align: center;
        color: #6c757d;
        font-size: 1.4rem;
        margin-bottom: 3rem;
        font-weight: 400;
        letter-spacing: 0.3px;
    }
    
    .feature-card {
        background: linear-gradient(135deg, #ffffff 0%, #f8f9fa 100%);
        padding: 2.5rem;
        border-radius: 25px;
        box-shadow: 0 15px 35px rgba(0,0,0,0.08);
        border: 1px solid #f0f0f0;
        transition: all 0.4s cubic-bezier(0.175, 0.885, 0.32, 1.275);
        margin-bottom: 2rem;
        height: 100%;
        cursor: pointer;
        position: relative;
        overflow: hidden;
    }
    
    .feature-card::before {
        content: '';
        position: absolute;
        top: 0;
        left: -100%;
        width: 100%;
        height: 100%;
        background: linear-gradient(90deg, transparent, rgba(255,255,255,0.4), transparent);
        transition: left 0.6s;
    }
    
    .feature-card:hover::before {
        left: 100%;
    }
    
    .feature-card:hover {
        transform: translateY(-10px) scale(1.02);
        box-shadow: 0 25px 50px rgba(102, 126, 234, 0.15);
        border-color: #667eea;
    }
    
    .feature-icon {
        font-size: 4.5rem;
        margin-bottom: 1.5rem;
        text-align: center;
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        filter: drop-shadow(0 4px 8px rgba(102, 126, 234, 0.3));
    }
    
    .legal-result {
        background: linear-gradient(135deg, #f8f9fa 0%, #e9ecef 100%);
        padding: 2.5rem;
        border-radius: 20px;
        border-left: 6px solid #667eea;
        margin: 1.5rem 0;
        box-shadow: 0 10px 30px rgba(0,0,0,0.08);
        line-height: 1.8;
        white-space: pre-line;
        border: 1px solid #e9ecef;
        animation: fadeIn 0.6s ease-in;
    }
    
    .simplified-text {
        background: linear-gradient(135deg, #e3f2fd 0%, #f3f9ff 100%);
        padding: 2.2rem;
        border-radius: 18px;
        border-left: 6px solid #2196f3;
        margin: 1.5rem 0;
        box-shadow: 0 10px 30px rgba(33, 150, 243, 0.1);
        line-height: 1.8;
        white-space: pre-line;
        border: 1px solid #e1f5fe;
        animation: fadeIn 0.6s ease-in;
    }
    
    .quick-feature {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        padding: 2.2rem 1.5rem;
        border-radius: 20px;
        text-align: center;
        margin: 0.8rem;
        cursor: pointer;
        transition: all 0.4s cubic-bezier(0.175, 0.885, 0.32, 1.275);
        border: none;
        box-shadow: 0 12px 30px rgba(102, 126, 234, 0.25);
        position: relative;
        overflow: hidden;
    }
    
    .quick-feature::before {
        content: '';
        position: absolute;
        top: 0;
        left: -100%;
        width: 100%;
        height: 100%;
        background: linear-gradient(90deg, transparent, rgba(255,255,255,0.2), transparent);
        transition: left 0.6s;
    }
    
    .quick-feature:hover::before {
        left: 100%;
    }
    
    .quick-feature:hover {
        transform: translateY(-8px) scale(1.05);
        box-shadow: 0 25px 45px rgba(102, 126, 234, 0.35);
    }
    
    .legal-tip {
        background: linear-gradient(135deg, #fff3e0 0%, #ffecb3 100%);
        padding: 1.5rem;
        border-radius: 15px;
        margin: 1rem 0;
        border-left: 5px solid #ff9800;
        box-shadow: 0 6px 20px rgba(255, 152, 0, 0.1);
        transition: transform 0.3s ease;
        border: 1px solid #ffe0b2;
    }
    
    .legal-tip:hover {
        transform: translateX(5px);
    }
    
    .emergency-box {
        background: linear-gradient(135deg, #ff6b6b 0%, #ff8e8e 100%);
        color: white;
        padding: 1.8rem;
        border-radius: 15px;
        margin: 1rem 0;
        text-align: center;
        box-shadow: 0 10px 25px rgba(255, 107, 107, 0.25);
        transition: all 0.3s ease;
        border: none;
        position: relative;
        overflow: hidden;
    }
    
    .emergency-box::before {
        content: '';
        position: absolute;
        top: -50%;
        left: -50%;
        width: 200%;
        height: 200%;
        background: radial-gradient(circle, rgba(255,255,255,0.1) 0%, transparent 70%);
        animation: pulse 2s infinite;
    }
    
    .emergency-box:hover {
        transform: scale(1.05);
        box-shadow: 0 15px 35px rgba(255, 107, 107, 0.4);
    }
    
    .language-selector {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        padding: 0.8rem 1.5rem;
        border-radius: 25px;
        border: none;
        box-shadow: 0 6px 20px rgba(102, 126, 234, 0.3);
        font-weight: 500;
    }
    
    .stButton>button {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        border: none;
        padding: 1.2rem 2.5rem;
        border-radius: 50px;
        font-weight: 600;
        transition: all 0.4s cubic-bezier(0.175, 0.885, 0.32, 1.275);
        width: 100%;
        box-shadow: 0 10px 25px rgba(102, 126, 234, 0.3);
        position: relative;
        overflow: hidden;
    }
    
    .stButton>button::before {
        content: '';
        position: absolute;
        top: 0;
        left: -100%;
        width: 100%;
        height: 100%;
        background: linear-gradient(90deg, transparent, rgba(255,255,255,0.3), transparent);
        transition: left 0.5s;
    }
    
    .stButton>button:hover::before {
        left: 100%;
    }
    
    .stButton>button:hover {
        transform: translateY(-4px);
        box-shadow: 0 18px 40px rgba(102, 126, 234, 0.4);
    }
    
    .stats-card {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        padding: 2.2rem;
        border-radius: 20px;
        text-align: center;
        margin: 1rem;
        box-shadow: 0 18px 35px rgba(102, 126, 234, 0.25);
        transition: transform 0.3s ease;
        position: relative;
        overflow: hidden;
    }
    
    .stats-card::before {
        content: '';
        position: absolute;
        top: -50%;
        left: -50%;
        width: 200%;
        height: 200%;
        background: radial-gradient(circle, rgba(255,255,255,0.1) 0%, transparent 70%);
    }
    
    .stats-card:hover {
        transform: translateY(-5px);
    }
    
    .language-badge {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        padding: 0.6rem 1.4rem;
        border-radius: 25px;
        font-size: 0.9rem;
        margin-left: 0.5rem;
        box-shadow: 0 4px 15px rgba(102, 126, 234, 0.3);
        font-weight: 500;
    }
    
    @keyframes fadeIn {
        from { opacity: 0; transform: translateY(20px); }
        to { opacity: 1; transform: translateY(0); }
    }
    
    @keyframes pulse {
        0% { transform: scale(0.5); opacity: 0; }
        50% { opacity: 0.3; }
        100% { transform: scale(1.2); opacity: 0; }
    }
    
    .section-divider {
        height: 3px;
        background: linear-gradient(90deg, transparent, #667eea, transparent);
        margin: 3rem 0;
        border: none;
    }
    
    .floating-animation {
        animation: float 6s ease-in-out infinite;
    }
    
    @keyframes float {
        0% { transform: translateY(0px); }
        50% { transform: translateY(-10px); }
        100% { transform: translateY(0px); }
    }
</style>
""", unsafe_allow_html=True)

# Initialize translator
translator = Translator()

# Language options
LANGUAGE_OPTIONS = {
    'English': 'en',
    'Hindi': 'hi', 
    'Tamil': 'ta',
    'Telugu': 'te',
    'Bengali': 'bn',
    'Marathi': 'mr',
    'Gujarati': 'gu',
    'Kannada': 'kn',
    'Malayalam': 'ml',
    'Punjabi': 'pa',
    'Urdu': 'ur'
}

# Get API Keys
GROQ_API_KEY = os.getenv('GROQ_API_KEY')

def init_session_state():
    """Initialize session state variables"""
    if 'complaint_result' not in st.session_state:
        st.session_state.complaint_result = None
    if 'simplified_text' not in st.session_state:
        st.session_state.simplified_text = None
    if 'current_feature' not in st.session_state:
        st.session_state.current_feature = "home"
    if 'selected_language' not in st.session_state:
        st.session_state.selected_language = 'English'

def extract_text_from_file(uploaded_file):
    """Extract text from uploaded files"""
    try:
        if uploaded_file.type == "application/pdf":
            pdf_reader = PyPDF2.PdfReader(uploaded_file)
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text()
            return text
        
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = docx.Document(uploaded_file)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        
        elif uploaded_file.type == "text/plain":
            return uploaded_file.read().decode("utf-8")
        
        else:
            return None
    except Exception as e:
        return f"Error reading file: {str(e)}"

def analyze_with_groq(prompt):
    """Use Groq API for analysis"""
    try:
        url = "https://api.groq.com/openai/v1/chat/completions"
        headers = {
            'Authorization': f'Bearer {GROQ_API_KEY}',
            'Content-Type': 'application/json'
        }
        data = {
            "messages": [
                {
                    "role": "system", 
                    "content": "You are a legal expert specializing in Indian law. Provide accurate, specific legal information with exact law sections and practical guidance. Use simple, easy-to-understand language suitable for common people."
                },
                {
                    "role": "user", 
                    "content": prompt
                }
            ],
            "model": "llama-3.1-8b-instant",
            "max_tokens": 2000,
            "temperature": 0.3
        }
        
        response = requests.post(url, headers=headers, json=data, timeout=30)
        
        if response.status_code == 200:
            result = response.json()
            return result['choices'][0]['message']['content']
        else:
            return f"❌ Groq API error: {response.status_code}"
    except Exception as e:
        return f"❌ Groq API failed: {str(e)}"

def analyze_with_ai(complaint_text, analysis_type="complaint"):
    """AI analysis using Groq API with simple language prompts"""
    
    if analysis_type == "complaint":
        prompt = f"""
        As an Indian legal expert, analyze this complaint in very simple, easy-to-understand language:

        USER COMPLAINT: {complaint_text}

        Provide analysis in this clear format:

        📋 APPLICABLE LAWS
        • List Indian laws and sections in simple terms
        • Explain what each law means for this case
        • Use everyday language anyone can understand

        🛡️ YOUR LEGAL RIGHTS  
        • Explain rights in simple, practical terms
        • Mention time limits clearly
        • Tell what protections you have

        ⚖️ POSSIBLE SOLUTIONS
        • List practical steps to take
        • Explain legal remedies simply
        • Mention compensation possibilities

        🚀 IMMEDIATE ACTIONS
        • Step-by-step what to do now
        • Documents needed
        • Where to go for help

        Use very simple words. Avoid complex legal terms. Explain like you're helping a friend.
        Make it easy for anyone to understand their legal options.
        """
    else:
        prompt = f"""
        Explain this legal document in very simple, everyday language:

        DOCUMENT TEXT: {complaint_text}

        Break it down like this:

        📄 WHAT THIS DOCUMENT IS ABOUT
        • Simple summary in plain English
        • Main purpose explained clearly
        • Who is involved

        🔍 KEY POINTS MADE SIMPLE
        • Explain important terms in easy words
        • What each party needs to do
        • Main responsibilities

        ⏰ IMPORTANT DATES & DEADLINES
        • When things need to happen
        • Time limits to remember
        • Critical dates

        ⚠️ THINGS TO WATCH OUT FOR
        • Potential risks explained simply
        • Important warnings
        • What could go wrong

        Use extremely simple language. No legal jargon. Explain like you're talking to someone with no legal background.
        Make sure every point is crystal clear and easy to understand.
        """
    
    result = analyze_with_groq(prompt)
    return result

def generate_ai_template(template_type, user_input):
    """Generate AI-powered legal template in simple language"""
    prompt = f"""
    Create a {template_type} based on these details: {user_input}
    
    Make it:
    • Very simple and easy to understand
    • Use clear, everyday language
    • Include helpful instructions
    • Ready to use with [placeholders]
    • Suitable for people without legal knowledge
    
    Keep it practical and user-friendly.
    """
    
    result = analyze_with_groq(prompt)
    return result

def translate_text(text, target_lang):
    """Translate text to target language with better formatting"""
    try:
        if target_lang != 'English':
            # First, simplify the English text for better translation
            simplified_prompt = f"Simplify this legal text into very clear, simple English that will be easy to translate: {text}"
            simplified_text = analyze_with_groq(simplified_prompt)
            
            # Then translate the simplified version
            translated = translator.translate(simplified_text, dest=LANGUAGE_OPTIONS[target_lang])
            return translated.text
        return text
    except Exception as e:
        return f"Translation note: For best understanding, please consult the English version.\n\n{text}"

def create_hero_section():
    """Create the main hero section with clean design"""
    col1, col2, col3 = st.columns([1, 2, 1])
    
    with col2:
        st.markdown('<div class="main-header floating-animation">⚖️ LawLens</div>', unsafe_allow_html=True)
        st.markdown('<div class="tagline">Your Simple Legal Guide • Clear Advice • Easy to Understand</div>', unsafe_allow_html=True)
        
        # Clean language selector - NO STATUS BADGE
        lang_col1, lang_col2, lang_col3 = st.columns([1, 2, 1])
        with lang_col2:
            selected_lang = st.selectbox(
                "🌍 Choose Your Language",
                options=list(LANGUAGE_OPTIONS.keys()),
                key="main_lang"
            )
            st.session_state.selected_language = selected_lang

def create_quick_features():
    """Create quick access features with better design"""
    st.markdown("## 🚀 Quick Legal Help")
    
    col1, col2, col3, col4 = st.columns(4)
    
    with col1:
        if st.button("📝 **Case Analysis**\n\nUnderstand your legal situation", use_container_width=True):
            st.session_state.current_feature = "complaint"
    
    with col2:
        if st.button("📄 **Document Help**\n\nSimplify legal papers", use_container_width=True):
            st.session_state.current_feature = "document"
    
    with col3:
        if st.button("⚡ **Ready Templates**\n\nEasy-to-use formats", use_container_width=True):
            st.session_state.current_feature = "templates"
    
    with col4:
        if st.button("🆘 **Emergency**\n\nImmediate help contacts", use_container_width=True):
            st.session_state.current_feature = "emergency"

def create_stats_section():
    """Create statistics section with better design"""
    st.markdown("## ✨ Why Choose LawLens")
    
    col1, col2, col3, col4 = st.columns(4)
    
    with col1:
        st.markdown('<div class="stats-card"><h3>🤖 AI Powered</h3><p>Smart Legal Analysis</p></div>', unsafe_allow_html=True)
    
    with col2:
        st.markdown('<div class="stats-card"><h3>🌍 Multi-Language</h3><p>11 Indian Languages</p></div>', unsafe_allow_html=True)
    
    with col3:
        st.markdown('<div class="stats-card"><h3>⚡ Instant</h3><p>Quick Responses</p></div>', unsafe_allow_html=True)
    
    with col4:
        st.markdown('<div class="stats-card"><h3>🆓 Free</h3><p>No Cost Help</p></div>', unsafe_allow_html=True)

def complaint_analyzer():
    """Complaint analyzer feature with simple language"""
    st.markdown("## 📝 Understand Your Legal Case")
    
    complaint = st.text_area(
        "**Tell us what happened in your own words:**",
        placeholder="Example: I bought a phone that stopped working after 2 days. The shop owner is refusing to help me...",
        height=120
    )
    
    if st.button("🔍 Get Simple Legal Advice", use_container_width=True):
        if complaint:
            with st.spinner("🤖 Understanding your situation... Making it simple..."):
                result = analyze_with_ai(complaint, "complaint")
                
                # Enhanced translation for better readability
                if st.session_state.selected_language != 'English':
                    result = translate_text(result, st.session_state.selected_language)
                
                st.success("✅ Clear Legal Guidance Ready!")
                st.markdown(f'<div class="legal-result">{result}</div>', unsafe_allow_html=True)
        else:
            st.warning("Please tell us about your situation")

def document_simplifier():
    """Document simplifier feature with simple language"""
    st.markdown("## 📄 Make Legal Papers Simple")
    
    # File upload
    uploaded_file = st.file_uploader(
        "**Upload your legal document:**",
        type=['pdf', 'docx', 'txt'],
        help="We support PDF, Word documents, and text files"
    )
    
    legal_text = st.text_area(
        "**Or paste the document text here:**",
        placeholder="Paste any legal document, contract, or notice you want to understand better...",
        height=200
    )
    
    if uploaded_file:
        extracted_text = extract_text_from_file(uploaded_file)
        if extracted_text and not extracted_text.startswith("Error"):
            legal_text = extracted_text
            st.success(f"📁 Document loaded successfully!")
        elif extracted_text.startswith("Error"):
            st.error(extracted_text)
    
    if st.button("✨ Make This Document Simple", use_container_width=True):
        if legal_text:
            with st.spinner("🤖 Reading your document... Making it easy to understand..."):
                result = analyze_with_ai(legal_text, "document")
                
                # Enhanced translation for better readability
                if st.session_state.selected_language != 'English':
                    result = translate_text(result, st.session_state.selected_language)
                
                st.success("✅ Document Made Simple!")
                st.markdown(f'<div class="simplified-text">{result}</div>', unsafe_allow_html=True)
        else:
            st.warning("Please provide a document to simplify")

def legal_templates():
    """Legal templates feature with simple language"""
    st.markdown("## 📋 Easy Legal Templates")
    
    template_options = {
        "📝 Legal Notice": "Send a formal legal notice",
        "🛒 Consumer Complaint": "Complain about products or services", 
        "🏠 Rent Agreement": "Create rental agreement",
        "💼 Work Contract": "Employment agreement",
        "🤝 Business Agreement": "Partnership contract",
        "💰 Loan Agreement": "Personal loan document"
    }
    
    template_choice = st.selectbox("Choose Template Type", list(template_options.keys()))
    
    st.markdown("### 📝 Tell Us What You Need")
    user_input = st.text_area(
        "Describe what you want in the template:",
        placeholder="Example: I need to send a notice to my landlord about security deposit return...",
        height=100
    )
    
    if st.button("🚀 Create Easy Template", use_container_width=True):
        if user_input:
            with st.spinner("🤖 Creating your simple template..."):
                template = generate_ai_template(template_choice, user_input)
                
                # Enhanced translation for better readability
                if st.session_state.selected_language != 'English':
                    template = translate_text(template, st.session_state.selected_language)
                
                st.success("✅ Your Template is Ready!")
                st.markdown(f'<div class="simplified-text">{template}</div>', unsafe_allow_html=True)
        else:
            st.warning("Please tell us what you need in the template")

def emergency_help():
    """Emergency help feature with better design"""
    st.markdown("## 🆘 Immediate Help Contacts")
    
    st.info("""
    **Important**: For urgent legal emergencies, contact these authorities directly. 
    This information is for immediate assistance when you need help right away.
    """)
    
    col1, col2 = st.columns(2)
    
    with col1:
        st.markdown('<div class="emergency-box">🚨 Police Emergency<br><strong>100</strong></div>', unsafe_allow_html=True)
        st.markdown('<div class="emergency-box">🚑 Medical Emergency<br><strong>108</strong></div>', unsafe_allow_html=True)
        st.markdown('<div class="emergency-box">👮 Women\'s Safety<br><strong>1091</strong></div>', unsafe_allow_html=True)
    
    with col2:
        st.markdown('<div class="emergency-box">👵 Senior Citizens<br><strong>14567</strong></div>', unsafe_allow_html=True)
        st.markdown('<div class="emergency-box">👶 Child Protection<br><strong>1098</strong></div>', unsafe_allow_html=True)
        st.markdown('<div class="emergency-box">📞 All Emergencies<br><strong>112</strong></div>', unsafe_allow_html=True)
    
    st.markdown("### 🏛️ Free Legal Help Centers")
    st.write("""
    • **District Legal Services** - Free legal aid in every district
    • **State Legal Authority** - State-level legal assistance  
    • **Local Legal Committees** - Help in your local area
    • **High Court Legal Help** - Court-level assistance
    • **National Legal Helpline** - 011-2338-2743
    """)

def create_legal_tips():
    """Create legal tips section with better design"""
    st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
    st.markdown("## 💡 Smart Legal Tips")
    
    tips = [
        "⚖️ Always talk to a qualified lawyer for serious legal matters",
        "📞 Save emergency numbers in your phone for quick access",
        "⏰ Know time limits for different legal complaints",
        "📧 Keep copies of all important messages and documents",
        "🔍 Check lawyer credentials before hiring anyone",
        "💰 Save financial records for at least 3 years",
        "📝 Read everything carefully before signing",
        "🤝 Try friendly solutions before going to court",
        "🏛️ Free legal help is available if you qualify",
        "📱 Use safe methods for important legal messages"
    ]
    
    # Display tips in two columns with better design
    col1, col2 = st.columns(2)
    
    for i, tip in enumerate(tips):
        with col1 if i % 2 == 0 else col2:
            st.markdown(f'<div class="legal-tip">{tip}</div>', unsafe_allow_html=True)

def main():
    """Main application function"""
    try:
        init_session_state()
        
        # Hero Section - Clean without status badge
        create_hero_section()
        
        # Stats Section
        create_stats_section()
        
        # Quick Features
        create_quick_features()
        
        st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
        
        # Main Features based on selection
        if st.session_state.current_feature == "complaint":
            complaint_analyzer()
        elif st.session_state.current_feature == "document":
            document_simplifier()
        elif st.session_state.current_feature == "templates":
            legal_templates()
        elif st.session_state.current_feature == "emergency":
            emergency_help()
        else:
            # Default home view with better design
            st.markdown("## 🎯 How We Help You Understand Law")
            
            col1, col2, col3 = st.columns(3)
            
            with col1:
                st.markdown('<div class="feature-card"><div class="feature-icon">⚖️</div><h3>Case Analysis</h3><p>• Understand your legal situation<br>• Know your rights clearly<br>• Get practical next steps<br>• Simple explanations</p></div>', unsafe_allow_html=True)
            
            with col2:
                st.markdown('<div class="feature-card"><div class="feature-icon">📚</div><h3>Document Help</h3><p>• Make legal papers easy to read<br>• Understand contracts clearly<br>• Know what you\'re signing<br>• Spot important points</p></div>', unsafe_allow_html=True)
            
            with col3:
                st.markdown('<div class="feature-card"><div class="feature-icon">🛡️</div><h3>Legal Protection</h3><p>• Ready-to-use templates<br>• Emergency contacts<br>• Know your protections<br>• Preventive guidance</p></div>', unsafe_allow_html=True)
        
        # Legal Tips at the bottom
        create_legal_tips()
        
        # Clean Footer
        st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
        st.markdown(
            "<div style='text-align: center; color: #666; padding: 2rem;'>"
            f"⚖️ <strong>LawLens</strong> - Making Law Simple for Everyone • "
            f"<span class='language-badge'>{st.session_state.selected_language}</span><br>"
            "💡 Legal information provided is for guidance only. Always consult a qualified lawyer for important legal matters."
            "</div>", 
            unsafe_allow_html=True
        )
    
    except Exception as e:
        st.error(f"Application Error: {str(e)}")
        st.info("Please refresh the page and try again.")

if __name__ == "__main__":
    main()